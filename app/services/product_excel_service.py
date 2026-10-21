# -*- coding: utf-8 -*-
"""ساخت الگوی Excel و ورود امن گروهی کالاها."""

from __future__ import annotations

import logging
import math
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.config import CATEGORIES, PRODUCT_CODE_PADDING, PRODUCT_CODE_PREFIX, UNITS
from app.services.product_service import ProductService

logger = logging.getLogger(__name__)

HEADERS = ("کد کالا", "نام کالا", "واحد", "حداقل موجودی", "دسته", "قفسه")
DEFAULT_UNIT = "عدد"
DEFAULT_CATEGORY = "عمومی"
MAX_IMPORT_ROWS = 50_000  # سقف نرم برای جلوگیری از فایل‌های تصادفیِ بسیار بزرگ

# نوع خروجی preview: (rows قابل ورود، خطاها، تعداد ردیف غیرخالی، هشدارها)
PreviewResult = tuple[
    list[dict[str, Any]], list[str], int, list[str]
]

# --------------------------------------------------------------------- template


def create_product_template(path: str | Path) -> Path:
    """ساخت الگوی xlsx با راهنما و اعتبارسنجی سلول.

    - ردیف ۱: هدر (با کامنت راهنما روی A1)
    - ردیف ۲: خالی، برای شروع ورود کاربر
    - اعتبارسنجی لیست/عدد تا ردیف ۱۰۰۰
    - شیت «راهنما» توضیحات کامل
    """
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "فهرست کالاها"
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F1000"
    sheet.row_dimensions[1].height = 28

    for cell in sheet[1]:
        cell.font = Font(name="Tahoma", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102A43")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet["A1"].comment = Comment(
        "این ردیف هدر است. اطلاعات کالاها را از ردیف ۲ به بعد وارد کنید.",
        "انبارداری آسان",
    )
    sheet["D1"].comment = Comment(
        "عدد صفر یا مثبت. اگر خالی بماند، صفر در نظر گرفته می‌شود.",
        "انبارداری آسان",
    )

    for col, width in {"A": 20, "B": 30, "C": 16, "D": 20, "E": 18, "F": 20}.items():
        sheet.column_dimensions[col].width = width

    # ردیف ۲ خالی است تا کاربر از همان‌جا شروع کند (کامنت روی A1 است).
    sheet.append([None, None, None, None, None, None])

    units = DataValidation(
        type="list", formula1='"' + ",".join(UNITS) + '"', allow_blank=True
    )
    categories = DataValidation(
        type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True
    )
    minimum = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True
    )
    for validation, message, cell_range in (
        (units, "واحد را از فهرست انتخاب کنید.", "C2:C1000"),
        (categories, "دسته را از فهرست انتخاب کنید.", "E2:E1000"),
        (minimum, "حداقل موجودی باید صفر یا بیشتر باشد.", "D2:D1000"),
    ):
        validation.error = message
        sheet.add_data_validation(validation)
        validation.add(cell_range)

    guide = workbook.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 90
    guide.append(["ستون", "راهنما"])
    guidance = (
        ("کد کالا", "اجباری؛ توسط کاربر تعیین می‌شود. تکراری بودن کد فقط هشدار است."),
        ("نام کالا", "اجباری؛ نام تکراری فقط هشدار است و با تأیید کاربر قابل ورود است."),
        ("واحد", "اختیاری؛ اگر خالی باشد «عدد» ثبت می‌شود."),
        ("حداقل موجودی", "اختیاری؛ عدد صفر یا مثبت. پیش‌فرض صفر."),
        ("دسته", "اختیاری؛ از فهرست انتخاب شود. پیش‌فرض عمومی."),
        ("قفسه", "اختیاری؛ محل نگهداری کالا."),
        ("شناسه داخلی", "برنامه شناسه مستقل می‌سازد تا گردش‌ها با کد تکراری هم درست بمانند."),
        ("موجودی اولیه", "از فایل وارد نمی‌شود و صفر است؛ گردش ورود را از فرم اصلی ثبت کنید."),
        ("ورود گروهی", "تکرار کد یا نام هشدار دارد، اما ردیف معتبر پس از تأیید کاربر وارد می‌شود."),
        ("نوع سلول", "سلول‌های متنی را به‌صورت «متن» (Text) ذخیره کنید تا صفرهای ابتدایی حفظ شوند."),
    )
    for row in guidance:
        guide.append(row)

    for cell in guide[1]:
        cell.font = Font(name="Tahoma", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102A43")
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            if cell.row > 1:
                cell.font = Font(name="Tahoma", size=10, bold=(cell.column == 1))
    guide.freeze_panes = "A2"

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


# -------------------------------------------------------------------- parsing


def _text(value: Any) -> str:
    """تبدیل ایمن سلول Excel به رشتهٔ تمیز."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # جلوگیری از «123.0» برای کدهای عددی
        return str(int(value))
    return str(value).strip()


def _parse_min_stock(raw: Any) -> Optional[float]:
    """تبدیل ایمن ورودی به عدد متناهی غیرمنفی یا ``None``."""
    if raw is None or _text(raw) == "":
        return 0.0
    try:
        # پذیرش «1,5» (ویرگول فارسی/اروپایی) به‌عنوان جداکنندهٔ اعشار
        value = float(str(raw).strip().replace("،", ".").replace(",", "."))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(value) or value < 0:
        return None
    return value


def _locate_headers(header_row: tuple[Any, ...] | None) -> dict[str, int]:
    """نقشهٔ نام ستون → اندیس، با خطای مشخص در نبود ستون‌های اجباری."""
    normalized = [_text(value) for value in (header_row or ())]
    positions = {name: normalized.index(name) for name in HEADERS if name in normalized}
    if "نام کالا" not in positions:
        raise ValueError(
            "ستون «نام کالا» در ردیف اول فایل پیدا نشد. از الگوی رسمی برنامه استفاده کنید."
        )
    return positions


# --------------------------------------------------------------- preview


def preview_product_import(
    path: str | Path, service: ProductService
) -> PreviewResult:
    """خواندن فایل Excel و اعتبارسنجی ردیف‌ها بدون نوشتن در دیتابیس.

    خروجی:
    - ``valid``: ردیف‌های قابل درج (dict).
    - ``errors``: پیام‌های خطا با شمارهٔ ردیف Excel.
    - ``total_rows``: تعداد ردیف‌های غیرخالی.
    - ``warnings``: هشدارهای تکرار کد/نام (مانع درج نیستند).
    """
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"فایل Excel قابل خواندن نیست: {exc}") from exc

    try:
        sheet = (
            workbook["فهرست کالاها"]
            if "فهرست کالاها" in workbook.sheetnames
            else workbook.active
        )
        iterator = sheet.iter_rows(values_only=True)
        positions = _locate_headers(next(iterator, None))

        existing_products = service.all()
        existing_names = {service.normalize_name(p.name) for p in existing_products}
        existing_codes = {
            (p.user_code or "").casefold()
            for p in existing_products
            if p.user_code
        }

        accepted_names: set[str] = set()
        accepted_codes: set[str] = set()
        valid: list[dict[str, Any]] = []
        errors: list[str] = []
        warnings: list[str] = []
        total_rows = 0

        for excel_row, values in enumerate(iterator, start=2):
            values = values or ()
            if not any(_text(value) for value in values):
                continue
            total_rows += 1
            if total_rows > MAX_IMPORT_ROWS:
                errors.append(
                    f"ردیف {excel_row}: سقف مجاز ورود ({MAX_IMPORT_ROWS} ردیف) رد شد."
                )
                break

            def get(column: str, default: Any = "") -> Any:
                index = positions.get(column)
                return (
                    values[index]
                    if index is not None and index < len(values)
                    else default
                )

            user_code = _text(get("کد کالا"))
            name = _text(get("نام کالا"))
            unit = _text(get("واحد", DEFAULT_UNIT)) or DEFAULT_UNIT
            category = _text(get("دسته", DEFAULT_CATEGORY)) or DEFAULT_CATEGORY
            shelf = _text(get("قفسه"))

            if not user_code:
                errors.append(f"ردیف {excel_row}: کد کالا خالی است.")
                continue
            if not name:
                errors.append(f"ردیف {excel_row}: نام کالا خالی است.")
                continue

            minimum = _parse_min_stock(get("حداقل موجودی", 0))
            if minimum is None:
                errors.append(
                    f"ردیف {excel_row}: حداقل موجودی باید عددی متناهی و غیرمنفی باشد."
                )
                continue
            if unit not in UNITS:
                errors.append(
                    f"ردیف {excel_row}: واحد «{unit}» معتبر نیست؛ از فهرست الگو انتخاب کنید."
                )
                continue
            if category not in CATEGORIES:
                errors.append(
                    f"ردیف {excel_row}: دسته «{category}» معتبر نیست."
                )
                continue

            name_key = service.normalize_name(name)
            code_key = user_code.casefold()
            if name_key in existing_names or name_key in accepted_names:
                warnings.append(
                    f"ردیف {excel_row}: نام «{name}» تکراری است؛ با تأیید قابل ورود است."
                )
            if code_key in existing_codes or code_key in accepted_codes:
                warnings.append(
                    f"ردیف {excel_row}: کد «{user_code}» تکراری است؛ با تأیید قابل ورود است."
                )
            accepted_names.add(name_key)
            accepted_codes.add(code_key)
            valid.append({
                "user_code": user_code,
                "name": name,
                "unit": unit,
                "min_stock": minimum,
                "category": category,
                "shelf": shelf,
                "excel_row": excel_row,
            })

        return valid, errors, total_rows, warnings
    finally:
        workbook.close()


# -------------------------------------------------------------------- import


def _validate_import_row(item: dict[str, Any], index: int) -> dict[str, Any]:
    """اعتبارسنجی صریح یک ردیف قبل از درج (لایهٔ دفاعی در برابر دستکاری)."""
    user_code = str(item.get("user_code") or "").strip()
    name = str(item.get("name") or "").strip()
    if not user_code:
        raise ValueError(f"ردیف {index}: کد کالا خالی است.")
    if not name:
        raise ValueError(f"ردیف {index}: نام کالا خالی است.")

    unit = str(item.get("unit") or DEFAULT_UNIT).strip() or DEFAULT_UNIT
    category = str(item.get("category") or DEFAULT_CATEGORY).strip() or DEFAULT_CATEGORY
    if unit not in UNITS:
        raise ValueError(f"ردیف {index}: واحد «{unit}» معتبر نیست.")
    if category not in CATEGORIES:
        raise ValueError(f"ردیف {index}: دسته «{category}» معتبر نیست.")

    minimum = _parse_min_stock(item.get("min_stock", 0))
    if minimum is None:
        raise ValueError(
            f"ردیف {index}: حداقل موجودی باید عددی متناهی و غیرمنفی باشد."
        )

    return {
        "user_code": user_code,
        "name": name,
        "unit": unit,
        "category": category,
        "shelf": str(item.get("shelf") or "").strip(),
        "min_stock": minimum,
    }


def _next_product_number(existing_codes: list[str]) -> int:
    """بزرگ‌ترین شمارهٔ استفاده‌شده + ۱ (برای جلوگیری از استفادهٔ مجدد)."""
    prefix_len = len(PRODUCT_CODE_PREFIX)
    numbers: list[int] = []
    for code in existing_codes:
        if not code.startswith(PRODUCT_CODE_PREFIX):
            continue
        suffix = code[prefix_len:]
        if suffix.isdigit():
            numbers.append(int(suffix))
    return max(numbers, default=0) + 1


def _format_product_code(number: int) -> str:
    return f"{PRODUCT_CODE_PREFIX}{number:0{PRODUCT_CODE_PADDING}d}"


def import_product_rows(
    service: ProductService, rows: list[dict[str, Any]]
) -> int:
    """درج اتمیک ردیف‌های اعتبارسنجی‌شده و تولید شناسهٔ داخلی یکتا.

    در صورت خطا در هر ردیف، کل تراکنش rollback می‌شود.
    """
    if not rows:
        return 0

    prepared = [
        _validate_import_row(item, index)
        for index, item in enumerate(rows, start=1)
    ]

    with service.db.transaction() as conn:
        existing_codes = [
            str(row["code"]) for row in conn.execute("SELECT code FROM products")
        ]
        next_number = _next_product_number(existing_codes)

        payload: list[tuple[str, str, str, str, float, str, str]] = []
        for row in prepared:
            internal_code = _format_product_code(next_number)
            next_number += 1
            payload.append((
                internal_code,
                row["user_code"],
                row["name"],
                row["unit"],
                row["min_stock"],
                row["category"],
                row["shelf"],
            ))

        conn.executemany(
            "INSERT INTO products "
            "(code, user_code, name, unit, min_stock, category, shelf) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            payload,
        )

    logger.info("Imported %d products from Excel", len(prepared))
    return len(prepared)


__all__ = [
    "HEADERS",
    "create_product_template",
    "import_product_rows",
    "preview_product_import",
]
