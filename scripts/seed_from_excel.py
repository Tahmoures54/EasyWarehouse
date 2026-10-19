# -*- coding: utf-8 -*-
"""
ورود داده از فایل اکسل انبار به EasyWarehouse
  python scripts/seed_from_excel.py [مسیر_اکسل]
"""
import sys
import sqlite3
from pathlib import Path
from datetime import date

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
DEFAULT_XLSX = Path("/home/workdir/attachments/anbar.xlsx")


def to_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return 0.0
    s = str(v).strip()
    if s in ("", "0"):
        return 0.0
    if s == "یک":
        return 1.0
    try:
        return float(s.replace(",", ""))
    except Exception:
        return 0.0


def calc_min(final, source):
    if source == "مصرفی":
        if final >= 100:
            return 20
        if final >= 30:
            return 10
        if final >= 10:
            return 5
        return max(1, int(final)) if final > 0 else 1
    if final >= 50:
        return 10
    if final >= 10:
        return 3
    return 1


def extract(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    items, seen = [], set()

    ws = wb["انبار لوازم مصرفی"]
    for row_idx in range(4, ws.max_row + 1):
        name = ws[f"E{row_idx}"].value
        if not name or not str(name).strip():
            continue
        name = str(name).strip()
        unit = str(ws[f"F{row_idx}"].value or "عدد").strip()
        shelf = ws[f"D{row_idx}"].value
        shelf = str(shelf).strip() if shelf else ""
        prev = to_num(ws[f"G{row_idx}"].value)
        tin = tout = tret = 0.0
        for day in range(31):
            base = 13 + day * 3
            tin += to_num(ws.cell(row=row_idx, column=base).value)
            tout += to_num(ws.cell(row=row_idx, column=base + 1).value)
            tret += to_num(ws.cell(row=row_idx, column=base + 2).value)
        final = max(0.0, prev + tin + tret - tout)
        key = ("مصرفی", name, unit)
        if key in seen:
            continue
        seen.add(key)
        items.append((name, unit, calc_min(final, "مصرفی"), "مصرفی", shelf, final))

    ws2 = wb["انبار قطعات و لوازم تکنیکال"]
    for row_idx in range(4, ws2.max_row + 1):
        name = ws2[f"E{row_idx}"].value
        if not name or not str(name).strip():
            continue
        name = str(name).strip()
        unit = str(ws2[f"F{row_idx}"].value or "عدد").strip()
        prev = to_num(ws2[f"G{row_idx}"].value)
        tin = tout = tret = 0.0
        for day in range(31):
            base = 12 + day * 3
            tin += to_num(ws2.cell(row=row_idx, column=base).value)
            tout += to_num(ws2.cell(row=row_idx, column=base + 1).value)
            tret += to_num(ws2.cell(row=row_idx, column=base + 2).value)
        final = max(0.0, prev + tin + tret - tout)
        key = ("تکنیکال", name, unit)
        if key in seen:
            continue
        seen.add(key)
        items.append((name, unit, calc_min(final, "تکنیکال"), "تکنیکال", "", final))
    return items


def seed(xlsx_path):
    db_path = ROOT / "data" / "anbar.db"
    db_path.parent.mkdir(parents=True, exist_ok=True)
    if db_path.exists():
        db_path.unlink()
        print("دیتابیس قبلی پاک شد.")

    conn = sqlite3.connect(str(db_path))
    conn.executescript((ROOT / "app" / "database" / "schema.sql").read_text(encoding="utf-8"))

    items = extract(xlsx_path)
    print(f"کالاهای استخراج‌شده: {len(items)}")

    today = date.today().isoformat()
    prods, txs = [], []
    for i, (name, unit, mn, cat, shelf, final) in enumerate(items, 1):
        code = f"P{i:03d}"
        prods.append((code, name, unit, mn, cat, shelf))
        if final > 0:
            txs.append((code, "IN", final, today, "موجودی اولیه از اکسل"))

    conn.executemany(
        "INSERT INTO products (code, name, unit, min_stock, category, shelf) VALUES (?,?,?,?,?,?)",
        prods,
    )
    conn.executemany(
        "INSERT INTO transactions (code, type, qty, date, note) VALUES (?,?,?,?,?)",
        txs,
    )
    conn.commit()
    conn.close()
    print(f"ثبت کالا: {len(prods)} | موجودی اولیه: {len(txs)}")
    print(f"فایل: {db_path}")


if __name__ == "__main__":
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"فایل یافت نشد: {xlsx}")
        sys.exit(1)
    print(f"خواندن: {xlsx}")
    seed(xlsx)
    print("تمام.")
